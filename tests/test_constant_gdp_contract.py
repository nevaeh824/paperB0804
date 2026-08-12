import csv
import importlib.util
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BUILDER_PATH = ROOT / "data0804" / "build_invest_panel_weo.py"
ANALYSIS_CSV = ROOT / "data0804" / "invest_panel_weo.csv"
WEO_XLSX = ROOT / "data0804" / "WEOApr2026all.xlsx"


def load_builder():
    spec = importlib.util.spec_from_file_location("build_invest_panel_weo", BUILDER_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


class ConstantGDPContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.builder = load_builder()

    def test_builder_reads_constant_gdp_from_weo(self):
        self.assertEqual(
            list(self.builder.WEO_FIELDS.items())[:3],
            [
                ("GGR_NGDP", "Revenue_gdp"),
                ("NGDP", "CurrentGDP"),
                ("NGDP_R", "ConstantGDP"),
            ],
        )

        weo_values, numeric_values, metadata, _ = self.builder.read_weo_values()
        self.assertIn(("AUS", 1995, "NGDP_R"), numeric_values)
        self.assertEqual(numeric_values[("AUS", 1995, "NGDP_R")], 1155.779)
        ngdp_r = metadata.loc[metadata["code"].eq("NGDP_R")]
        self.assertEqual(len(ngdp_r), 197)
        self.assertEqual(
            len({iso3 for iso3, _, code in numeric_values if code == "NGDP_R"}),
            196,
        )
        self.assertEqual(ngdp_r["unit"].drop_duplicates().tolist(), ["Domestic currency"])
        self.assertEqual(ngdp_r["scale"].drop_duplicates().tolist(), ["Billions"])
        self.builder.verify_weo_reconciliation(weo_values)

    def test_existing_constant_gdp_matches_weo_ngdp_r(self):
        workbook = load_workbook(
            WEO_XLSX,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        worksheet = workbook["Countries"]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        index = {name: position for position, name in enumerate(headers)}
        source = {}
        for row in rows:
            if row[index["INDICATOR.ID"]] != "NGDP_R":
                continue
            iso3 = str(row[index["COUNTRY.ID"]]).strip()
            for year in range(1995, 2024):
                value = row[index[str(year)]]
                if value not in (None, ""):
                    source[(iso3, year)] = float(value)
        workbook.close()

        output_rows = read_csv(ANALYSIS_CSV)
        output = {
            (row["iso3"], int(row["year"])): (
                None if row["ConstantGDP"] == "" else float(row["ConstantGDP"])
            )
            for row in output_rows
        }
        relevant_source = {key: source.get(key) for key in output}

        self.assertEqual(len(output), 1827)
        self.assertEqual(sum(value is not None for value in output.values()), 1825)
        self.assertEqual(
            {key for key, value in output.items() if value is None},
            {key for key, value in relevant_source.items() if value is None},
        )
        differences = [
            abs(value - relevant_source[key])
            for key, value in output.items()
            if value is not None
        ]
        self.assertEqual(max(differences), 0.0)
        self.assertEqual(sum(value <= 0 for value in output.values() if value is not None), 0)

    def test_generated_outputs_use_constant_gdp_control(self):
        baseline_rows = read_csv(ROOT / "baseline" / "stata_outputs" / "model_coefficients.csv")
        baseline_variables = {row["variable"] for row in baseline_rows}
        self.assertIn("ln_constantgdp", baseline_variables)
        self.assertNotIn("ln_currentgdp", baseline_variables)

        theta_rows = read_csv(ROOT / "empirical_theta" / "stata_outputs" / "model_coefficients.csv")
        spread_variables = {
            row["variable"] for row in theta_rows if row["model"] == "Spread_Interact_all"
        }
        tax_variables = {
            row["variable"] for row in theta_rows if row["model"].startswith("T")
        }
        self.assertIn("ln_constantgdp", spread_variables)
        self.assertNotIn("ln_currentgdp", spread_variables)
        self.assertTrue(
            {"CurrentGDP", "ConstantGDP", "ln_currentgdp", "ln_constantgdp"}.isdisjoint(
                tax_variables
            )
        )

        panel_header = set(
            read_csv(ROOT / "empirical_theta" / "stata_outputs" / "empirical_theta_panel.csv")[0]
        )
        self.assertIn("ConstantGDP", panel_header)
        self.assertIn("ln_constantgdp", panel_header)
        self.assertNotIn("ln_currentgdp", panel_header)

        doom_rows = read_csv(
            ROOT / "doomloop" / "stata_outputs" / "nostate_model_coefficients.csv"
        )
        doom_variables = {row["variable"] for row in doom_rows}
        self.assertTrue(
            {"CurrentGDP", "ConstantGDP", "ln_currentgdp", "ln_constantgdp"}.isdisjoint(
                doom_variables
            )
        )

        results = (ROOT / "paperB" / "paperB_results.md").read_text(encoding="utf-8")
        diagnostics = (ROOT / "paperB" / "paperB_diagnostics.md").read_text(
            encoding="utf-8"
        )
        self.assertIn(r"\ln(ConstantGDP)", results)
        self.assertNotIn(r"\ln(CurrentGDP)", results)
        self.assertIn("ln_constantgdp", diagnostics)
        self.assertNotIn("ln_currentgdp", diagnostics)


if __name__ == "__main__":
    unittest.main()
